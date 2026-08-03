import os
import json
import tempfile
import csv
import openpyxl
import logging
from django.http import HttpResponse
from django.conf import settings
from django.db import transaction
from rest_framework import viewsets, status
from rest_framework.response import Response
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.exceptions import ValidationError

from apps.imports.models import ImportSession, ReuploadSession, ImportRowResult, ImportWarning, ImportStatus
from apps.imports.serializers import ImportSessionSerializer, ReuploadSessionSerializer
from apps.imports.services import ImportService
from apps.tables.services import TableService
from apps.fields.services import FieldService
from apps.jobs.models import Job, JobType, JobStatus
from apps.jobs.tasks import run_job_task
from apps.tables.models import Table

logger = logging.getLogger(__name__)

class ImportSessionViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = ImportSessionSerializer

    def get_queryset(self):
        qs = ImportSession.objects.filter(organization_id=self.request.user.organization_id)
        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(table__name__icontains=search)
        status_param = self.request.query_params.get('status')
        if status_param:
            qs = qs.filter(status=status_param)
        return qs

    @action(detail=False, methods=['post'], url_path='preview')
    def preview(self, request):
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            raise ValidationError("excel_file is required.")
        
        fd, temp_path = tempfile.mkstemp(suffix='.xlsx')
        try:
            with os.fdopen(fd, 'wb') as tmp:
                for chunk in excel_file.chunks():
                    tmp.write(chunk)
            
            preview_data = ImportService.detect_columns_from_excel(temp_path)
            return Response(preview_data)
        finally:
            try:
                os.remove(temp_path)
            except Exception:
                pass

    @action(detail=False, methods=['post'], url_path='create-table')
    def create_table(self, request):
        excel_file = request.FILES.get('excel_file')
        table_name = request.data.get('table_name')
        columns_json = request.data.get('columns')
        
        if not excel_file or not table_name or not columns_json:
            raise ValidationError("excel_file, table_name, and columns are required.")
            
        try:
            columns = json.loads(columns_json)
        except Exception:
            raise ValidationError("columns must be valid JSON list.")
            
        user = request.user
        org_id = str(user.organization_id)
        
        with transaction.atomic():
            table = TableService.create_table(org_id, table_name, user)
            for col in columns:
                col_name = col.get('column_name')
                col_type = col.get('detected_type')
                col_req = col.get('is_required', False)
                col_uniq = col.get('is_unique', False)
                if col_name and col_type:
                    FieldService.create_field(
                        table_id=str(table.id),
                        name=col_name,
                        type=col_type,
                        is_unique=col_uniq,
                        is_required=col_req,
                        created_by=user
                    )
                    
        suffix = os.path.splitext(excel_file.name)[1] or '.xlsx'
        fd, temp_path = tempfile.mkstemp(suffix=suffix)
        with os.fdopen(fd, 'wb') as tmp:
            for chunk in excel_file.chunks():
                tmp.write(chunk)
                
        import_session = ImportSession.objects.create(
            user=user,
            organization_id=org_id,
            table=table,
            status=ImportStatus.PENDING
        )
        
        job = Job.objects.create(
            type=JobType.IMPORT_XLSX,
            status=JobStatus.QUEUED,
            payload={
                "import_session_id": str(import_session.id),
                "excel_path": temp_path,
                "table_id": str(table.id)
            },
            created_by=user
        )
        
        result = run_job_task.delay(str(job.id))
        job.celery_task_id = result.id
        job.save()
        
        return Response({
            "import_session": ImportSessionSerializer(import_session).data,
            "job_id": str(job.id)
        }, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='xlsx')
    def xlsx_import(self, request):
        table_id = request.data.get('table_id')
        excel_file = request.FILES.get('excel_file')
        zip_file = request.FILES.get('zip_file')
        
        if not table_id or not excel_file:
            raise ValidationError("table_id and excel_file are required.")
            
        table = Table.objects.filter(id=table_id, organization_id=request.user.organization_id).first()
        if not table:
            raise ValidationError("Table not found or not in your organization.")
            
        user = request.user
        org_id = str(user.organization_id)
        
        suffix_ex = os.path.splitext(excel_file.name)[1] or '.xlsx'
        fd_ex, temp_ex_path = tempfile.mkstemp(suffix=suffix_ex)
        with os.fdopen(fd_ex, 'wb') as tmp:
            for chunk in excel_file.chunks():
                tmp.write(chunk)
                
        temp_zip_path = None
        job_type = JobType.IMPORT_XLSX
        if zip_file:
            job_type = JobType.IMPORT_XLSX_ZIP
            suffix_zip = os.path.splitext(zip_file.name)[1] or '.zip'
            fd_zip, temp_zip_path = tempfile.mkstemp(suffix=suffix_zip)
            with os.fdopen(fd_zip, 'wb') as tmp:
                for chunk in zip_file.chunks():
                    tmp.write(chunk)
                    
        import_session = ImportSession.objects.create(
            user=user,
            organization_id=org_id,
            table=table,
            status=ImportStatus.PENDING
        )
        
        job = Job.objects.create(
            type=job_type,
            status=JobStatus.QUEUED,
            payload={
                "import_session_id": str(import_session.id),
                "excel_path": temp_ex_path,
                "zip_path": temp_zip_path,
                "table_id": str(table.id)
            },
            created_by=user
        )
        
        result = run_job_task.delay(str(job.id))
        job.celery_task_id = result.id
        job.save()
        
        return Response({
            "import_session": ImportSessionSerializer(import_session).data,
            "job_id": str(job.id)
        }, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['get'], url_path='error-report')
    def error_report(self, request, pk=None):
        import_session = self.get_object()
        report_format = request.query_params.get('report_format', 'xlsx').lower()
        
        results = import_session.row_results.all()
        warnings = import_session.warnings.all()
        
        warning_by_row = {}
        for w in warnings:
            if w.row_number not in warning_by_row:
                warning_by_row[w.row_number] = []
            warning_by_row[w.row_number].append(f"[{w.warning_type}] {w.message}")
            
        if report_format == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = f'attachment; filename="import_error_report_{import_session.id}.csv"'
            writer = csv.writer(response)
            writer.writerow(['Row Number', 'Status', 'Errors/Warnings', 'Row Data'])
            
            for res in results:
                row_errors = []
                if res.error_message:
                    row_errors.append(res.error_message)
                if res.row_number in warning_by_row:
                    row_errors.extend(warning_by_row[res.row_number])
                    
                writer.writerow([
                    res.row_number,
                    res.status,
                    "; ".join(row_errors),
                    json.dumps(res.row_data)
                ])
            return response
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Import Error Report"
            ws.append(['Row Number', 'Status', 'Errors/Warnings', 'Row Data'])
            
            for res in results:
                row_errors = []
                if res.error_message:
                    row_errors.append(res.error_message)
                if res.row_number in warning_by_row:
                    row_errors.extend(warning_by_row[res.row_number])
                    
                ws.append([
                    res.row_number,
                    res.status,
                    "; ".join(row_errors),
                    json.dumps(res.row_data)
                ])
                
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="import_error_report_{import_session.id}.xlsx"'
            wb.save(response)
            return response

class ReuploadSessionViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = ReuploadSessionSerializer

    def get_queryset(self):
        return ReuploadSession.objects.filter(organization_id=self.request.user.organization_id)

    @action(detail=False, methods=['post'], url_path='zip')
    def zip_reupload(self, request):
        table_id = request.data.get('table_id')
        zip_file = request.FILES.get('zip_file')
        
        if not table_id or not zip_file:
            raise ValidationError("table_id and zip_file are required.")
            
        table = Table.objects.filter(id=table_id, organization_id=request.user.organization_id).first()
        if not table:
            raise ValidationError("Table not found or not in your organization.")
            
        user = request.user
        org_id = str(user.organization_id)
        
        suffix = os.path.splitext(zip_file.name)[1] or '.zip'
        fd, temp_path = tempfile.mkstemp(suffix=suffix)
        with os.fdopen(fd, 'wb') as tmp:
            for chunk in zip_file.chunks():
                tmp.write(chunk)
                
        reupload_session = ReuploadSession.objects.create(
            user=user,
            organization_id=org_id,
            table=table,
            status=ImportStatus.PENDING
        )
        
        job = Job.objects.create(
            type=JobType.REUPLOAD_IMAGES,
            status=JobStatus.QUEUED,
            payload={
                "reupload_session_id": str(reupload_session.id),
                "zip_path": temp_path,
                "table_id": str(table.id)
            },
            created_by=user
        )
        
        result = run_job_task.delay(str(job.id))
        job.celery_task_id = result.id
        job.save()
        
        return Response({
            "reupload_session": ReuploadSessionSerializer(reupload_session).data,
            "job_id": str(job.id)
        }, status=status.HTTP_201_CREATED)
